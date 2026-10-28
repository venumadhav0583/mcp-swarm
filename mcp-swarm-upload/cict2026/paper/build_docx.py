"""Build CICT348.docx from main.tex using the official IEEE A4 Word template."""
import re, copy, subprocess, tempfile, os
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.oxml import OxmlElement, parse_xml
from lxml import etree

def inc(m): return open(m.group(1) + '.tex').read()
TEX = re.sub(r'\\input\{([^}]*)\}', inc, open('main.tex').read())
TPL = 'ieee-a4-template.docx'
bib_tex = TEX[TEX.index(r'\begin{thebibliography}{99}') + 27:TEX.index(r'\end{thebibliography}')]
bib = re.findall(r'\\bibitem\{([^}]*)\}(.*?)(?=\\bibitem|\Z)', bib_tex, flags=re.S)
num = {k: i + 1 for i, (k, _) in enumerate(bib)}
labels = {}
for kind, pre in (('figure\\*?', 'fig'), ('table', 'tab')):
    for i, m in enumerate(re.finditer(r'\\begin\{' + kind + r'\}.*?\\label\{([^}]*)\}', TEX, re.S)): pass
fi = ti = ei = 0
for m in re.finditer(r'\\begin\{(figure\*?|table|equation)\}.*?\\label\{([^}]*)\}', TEX, re.S):
    if m.group(1).startswith('figure'): fi += 1; labels[m.group(2)] = str(fi)
    elif m.group(1) == 'table': ti += 1; labels[m.group(2)] = ['I', 'II', 'III', 'IV', 'V'][ti - 1]
    else: ei += 1; labels[m.group(2)] = str(ei)
GREEK = {r'\alpha': 'α', r'\beta': 'β', r'\lambda': 'λ', r'\tau': 'τ', r'\pm': '±', r'\le': '≤', r'\geq': '≥', r'\times': '×', r'\propto': '∝'}
def math_txt(s):
    s = s.replace(r'\ ', ' ')
    for a, b in sorted(GREEK.items(), key=lambda x: -len(x[0])): s = s.replace(a, b)
    s = re.sub(r'_\{([^}]*)\}', r'\1', s); s = re.sub(r'_(\w)', r'\1', s); s = re.sub(r'\^\{([^}]*)\}', r'\1', s)
    s = s.replace('^', '').replace('{', '').replace('}', '').replace('<', '<').replace('-', '−')
    return s
def cites(m):
    ns = sorted(num[k.strip()] for k in m.group(1).split(','))
    return ', '.join(f'[{n}]' for n in ns)
def tex2txt(t):
    t = re.sub(r'\\cite\{([^}]*)\}', cites, t)
    t = re.sub(r'\\ref\{([^}]*)\}', lambda m: labels[m.group(1)], t)
    t = re.sub(r'\\url\{([^}]*)\}', r'\1', t)
    t = re.sub(r'\\texttt\{([^}]*)\}', lambda m: m.group(1).replace(r'\_', '_'), t)
    t = re.sub(r'\$([^$]*)\$', lambda m: math_txt(m.group(1)), t)
    for a, b in [(r'\_', '_'), ('~', '\u00a0'), (r'\,', '\u2009'), (r'\%', '%'), ('---', '\u2014'), ('--', '\u2013'),
                 ("``", '\u201c'), ("''", '\u201d'), (r'\&', '&'), (r'\ ', ' '), ('\\\\', ' '), (r'\textit{et al.}', 'et al.'), (r'\textbf{', ''), (r'\textit{', '')]:
        t = t.replace(a, b)
    t = t.replace("`", '\u2018').replace("'", '\u2019')
    return re.sub(r'\s+', ' ', t).strip()
def runs(p, t, size=None, bold=None):
    t = re.sub(r'\\cite\{([^}]*)\}', cites, t)
    parts = re.split(r'(\\textit\{(?:[^{}]|\{[^{}]*\})*\}|\$[^$]*\$)', t)
    for part in parts:
        if not part: continue
        if part.startswith('$') and part.endswith('$') and len(part) > 1:
            if re.search(r'[_^]', part):
                lead = ' ' if (p.runs and not p.runs[-1].text.endswith((' ', '(', '\u00a0', '-', '\u2013'))) else ''
                for k, (base, s1, s2) in enumerate(re.findall(r'([^_^]+)(?:[_^](?:\{([^}]*)\}|(\w)))?', part[1:-1])):
                    sub = s1 or s2
                    bt = math_txt(base)
                    if bt:
                        r = p.add_run((lead if k == 0 else '') + bt); r.italic = bt.strip().isalpha() or None
                        if size: r.font.size = Pt(size)
                        if bold: r.bold = True
                    if sub:
                        r = p.add_run(math_txt(sub)); r.font.subscript = True
                        if size: r.font.size = Pt(size)
                continue
        it = part.startswith(r'\textit{')
        raw = part[8:-1] if it else part
        txt = tex2txt(raw)
        if not txt: continue
        if raw[:1].isspace() or (not it and part[:1].isspace()): txt = ' ' + txt
        if part[-1:].isspace(): txt = txt + ' '
        r = p.add_run(txt); r.italic = it or None
        if bold: r.bold = True
        if size: r.font.size = Pt(size)
    prev = None
    for r in p.runs:
        if prev is not None and prev.text and r.text and not prev.text[-1].isspace() and r.text[0] not in ' .,;:)]\u2019' and prev.text[-1] not in '([\u201c\u00a0-':
            r.text = ' ' + r.text
        prev = r

doc = Document(TPL)
B = doc.element.body; final_sect = B.find(qn('w:sectPr'))
for el in list(B):
    if el is not final_sect: B.remove(el)
def para(style, text='', size=None):
    p = doc.add_paragraph(style=style)
    if text: runs(p, text, size)
    return p
title = tex2txt(re.search(r'\\title\{(.*?)\}\n', TEX, re.S).group(1))
para('paper title', title)
p = para('Author'); p.add_run('Venumadhav Kalluru').font.size = Pt(11)
for line in ['Independent Researcher', 'Cumming, Georgia, USA', 'greenkal2026@gmail.com']:
    p.add_run().add_break(); r = p.add_run(line); r.font.size = Pt(9)
    if line == 'Independent Researcher': r.italic = True
def set_cols(sectPr, n):
    cols = sectPr.find(qn('w:cols'))
    if cols is None: cols = OxmlElement('w:cols'); sectPr.append(cols)
    for k in list(cols.attrib): del cols.attrib[k]
    if n > 1: cols.set(qn('w:num'), str(n))
    cols.set(qn('w:space'), '360')
sec_p = doc.add_paragraph(style='Author'); sp = copy.deepcopy(final_sect); set_cols(sp, 1)
sec_p._p.get_or_add_pPr().append(sp)
t = final_sect.find(qn('w:type'))
if t is None: t = OxmlElement('w:type'); final_sect.insert(0, t)
t.set(qn('w:val'), 'continuous'); set_cols(final_sect, 2)
abs_t = re.search(r'\\begin\{abstract\}(.*?)\\end\{abstract\}', TEX, re.S).group(1)
p = para('Abstract'); r = p.add_run('Abstract\u2014'); r.bold = True; r.italic = True
runs(p, abs_t, bold=True)
kw = re.search(r'\\begin\{IEEEkeywords\}(.*?)\\end\{IEEEkeywords\}', TEX, re.S).group(1)
p = para('Keywords'); r = p.add_run('Keywords\u2014'); r.bold = True; r.italic = True
r = p.add_run(tex2txt(kw)); r.bold = True; r.italic = True

def border(cell, sides):
    tcPr = cell._tc.get_or_add_tcPr(); b = OxmlElement('w:tcBorders')
    for s, w in sides.items():
        e = OxmlElement(f'w:{s}'); e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), str(w)); e.set(qn('w:color'), '000000'); b.append(e)
    tcPr.append(b)
def add_table(block):
    cap = tex2txt(re.search(r'\\caption\{((?:[^{}]|\{[^{}]*\})*)\}', block).group(1))
    body = block[block.index(r'\toprule') + 8:block.index(r'\bottomrule')]
    note = re.search(r'\\multicolumn\{\d+\}\{[^}]*\}\{(.*)\}\s*\\end\{tabular\}', block, re.S)
    note = note.group(1).replace(r'\scriptsize', '') if note else None
    rows = []
    for ln in body.split('\\\\'):
        ln = ln.replace(r'\midrule', '').strip()
        if ln: rows.append([c.strip() for c in ln.split('&')])
    ncol = len(rows[0])
    L = [max(len(tex2txt(r[j])) for r in rows) + 2 for j in range(ncol)]
    widths = [8.45 * l / sum(L) for l in L]
    para('table head', cap).paragraph_format.keep_with_next = True
    tb = doc.add_table(rows=len(rows) + (1 if note else 0), cols=ncol)
    tb.alignment = WD_TABLE_ALIGNMENT.CENTER; tb.autofit = False
    tblPr = tb._tbl.tblPr; mar = OxmlElement('w:tblCellMar')
    for s in ('left', 'right'):
        e = OxmlElement(f'w:{s}'); e.set(qn('w:w'), '40'); e.set(qn('w:type'), 'dxa'); mar.append(e)
    tblPr.append(mar)
    for i, row in enumerate(rows):
        for j, txt in enumerate(row):
            c = tb.cell(i, j); c.width = Cm(widths[j])
            pp = c.paragraphs[0]; pp.style = doc.styles['table copy']
            pp.alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.CENTER
            runs(pp, txt)
            for r in pp.runs:
                r.font.size = Pt(7.5)
                if i == 0: r.bold = True
            sides = {'top': 8, 'bottom': 4} if i == 0 else {}
            if i == len(rows) - 1 and not note: sides['bottom'] = 8
            if sides: border(c, sides)
    if note:
        rr = tb.rows[-1]; c = rr.cells[0].merge(rr.cells[-1]); pp = c.paragraphs[0]; pp.style = doc.styles['table copy']
        runs(pp, note)
        for r in pp.runs: r.font.size = Pt(7)
        border(c, {'top': 8})
    grid = tb._tbl.tblGrid
    for gc, w in zip(grid.findall(qn('w:gridCol')), widths): gc.set(qn('w:w'), str(int(Cm(w).twips)))
    tw = OxmlElement('w:tblW'); tw.set(qn('w:w'), str(int(Cm(sum(widths)).twips))); tw.set(qn('w:type'), 'dxa')
    old = tblPr.find(qn('w:tblW'))
    if old is not None: tblPr.remove(old)
    tblPr.append(tw)
    for i, row in enumerate(tb.rows):
        trPr = row._tr.get_or_add_trPr(); trPr.append(OxmlElement('w:cantSplit'))
        for c in row.cells:
            for pp in c.paragraphs:
                pp.paragraph_format.keep_with_next = i < len(tb.rows) - 1; pp.paragraph_format.first_line_indent = 0
                pp.paragraph_format.space_before = Pt(1); pp.paragraph_format.space_after = Pt(1)
    sp_ = doc.add_paragraph(style='Body Text'); sp_.paragraph_format.space_after = Pt(0)
def figure(png, caption, width_cm=8.45):
    p = doc.add_paragraph(style='Body Text'); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = 0; p.paragraph_format.keep_with_next = True; p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    p.add_run().add_picture(png, width=Cm(width_cm))
    cp = para('figure caption'); nofmt(cp); cp.add_run('Fig. 2.\u2002'); runs(cp, caption.replace('(a) ', '(a)\u00a0'))
def wide_figure(png, caption):
    """Full-width figure between continuous section breaks (two-column text above and below)."""
    brk = doc.add_paragraph(style='Body Text'); brk.paragraph_format.space_after = Pt(0)
    sp2 = copy.deepcopy(final_sect); set_cols(sp2, 2); brk._p.get_or_add_pPr().append(sp2)
    p = doc.add_paragraph(style='Body Text'); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = 0; p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    p.paragraph_format.space_before = Pt(6)
    p.add_run().add_picture(png, width=Cm(16.2))
    cp = doc.add_paragraph(style='figure caption'); nofmt(cp); cp.add_run('Fig. 1.\u2002'); runs(cp, caption)
    cp.paragraph_format.space_after = Pt(8)
    sp1 = copy.deepcopy(final_sect); set_cols(sp1, 1); cp._p.get_or_add_pPr().append(sp1)
EQS = {}
def nofmt(p):
    pPr = p._p.get_or_add_pPr(); n = OxmlElement('w:numPr'); i = OxmlElement('w:numId'); i.set(qn('w:val'), '0'); n.append(i); pPr.append(n)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT if False else p.alignment
def omml(latex):
    with tempfile.TemporaryDirectory() as d:
        open(f'{d}/e.md', 'w').write('$' + latex + '$\n')
        subprocess.run(['pandoc', f'{d}/e.md', '-o', f'{d}/e.docx'], check=True)
        x = subprocess.run(['unzip', '-p', f'{d}/e.docx', 'word/document.xml'], capture_output=True).stdout
    root = etree.fromstring(x)
    return root.find('.//{http://schemas.openxmlformats.org/officeDocument/2006/math}oMath')
def equation(block):
    lab = re.search(r'\\label\{([^}]*)\}', block).group(1)
    latex = re.sub(r'\\label\{[^}]*\}', '', block[block.index('}') + 1:block.index(r'\end{equation}')]).strip()
    latex = latex.replace(r'\mathrm{sim}', r'\operatorname{sim}').replace(r'\,', ' ').replace(r'\qquad', r'\quad ').replace(r'\ ', ' ')
    p = doc.add_paragraph(style='equation')
    p.add_run('\t')
    p._p.append(omml(latex))
    p.add_run('\t(' + labels[lab] + ')')
FLOAT_FUN = {'fig:arch': lambda b: wide_figure('figs/fig1_arch.png', tex2txt(re.search(r'\\caption\{(.*?)\}\n', b, re.S).group(1))),
             'fig:council': lambda b: figure('figs/fig2_council.png', tex2txt(re.search(r'\\caption\{(.*?)\}\n', b, re.S).group(1)))}
main = TEX[TEX.index(r'\begin{figure*}'):TEX.index(r'\begin{thebibliography}')]
store = {}
def marker(m):
    k = f'@@F{len(store)}@@'; store[k] = m.group(0); return f'\n\n{k}\n\n'
main = re.sub(r'\\begin\{(table|figure\*?)\}.*?\\end\{\1\}', marker, main, flags=re.S)
main = re.sub(r'\n?\\begin\{equation\}.*?\\end\{equation\}\n?', marker, main, flags=re.S)
pending_wide = None
for block in re.split(r'\n\s*\n', main):
    block = block.strip()
    if not block: continue
    if block in store:
        b = store[block].strip()
        if b.startswith(r'\begin{equation}'): equation(b)
        elif b.startswith(r'\begin{table}'): add_table(b)
        else:
            lab = re.search(r'\\label\{([^}]*)\}', b).group(1)
            if lab == 'fig:arch': pending_wide = b
            else: FLOAT_FUN[lab](b)
        continue
    # equations embedded inside paragraphs: split on markers
    while True:
        m = re.match(r'\\(sub)?section\{(.*?)\}\s*(.*)$', block, re.S)
        if not m: break
        if m.group(2) == 'Methods' and pending_wide is not None:
            FLOAT_FUN['fig:arch'](pending_wide); pending_wide = None
        para('Heading 2' if m.group(1) else 'Heading 1', tex2txt(m.group(2))); block = m.group(3).strip()
        if m.group(2) == 'Related Works' and pending_wide is not None:
            pass
    if not block: continue
    for seg in re.split(r'(@@F\d+@@)', block):
        seg = seg.strip()
        if not seg: continue
        if seg in store: equation(store[seg].strip()); continue
        seg = re.sub(r'\\label\{[^}]*\}', '', seg)
        pp = para('Body Text', seg.replace('\n', ' '))
        if seg.startswith('enter the prompt') or seg.startswith('where ') or seg.startswith('is returned'):
            pp.paragraph_format.first_line_indent = 0
    # anchor the wide figure after the first Related Works paragraph
    if pending_wide is not None and block.startswith('Guo et al.'):
        FLOAT_FUN['fig:arch'](pending_wide); pending_wide = None
p = para('Heading 5', 'References')
for k, t in bib:
    p = doc.add_paragraph(style='references'); p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    runs(p, t.strip())
ORDER = {'pPr': ['pStyle','keepNext','keepLines','pageBreakBefore','framePr','widowControl','numPr','suppressLineNumbers','pBdr','shd','tabs','suppressAutoHyphens','kinsoku','wordWrap','overflowPunct','topLinePunct','autoSpaceDE','autoSpaceDN','bidi','adjustRightInd','snapToGrid','spacing','ind','contextualSpacing','mirrorIndents','suppressOverlap','jc','textDirection','textAlignment','textboxTightWrap','outlineLvl','divId','cnfStyle','rPr','sectPr','pPrChange'], 'tblPr': ['tblStyle','tblpPr','tblOverlap','bidiVisual','tblStyleRowBandSize','tblStyleColBandSize','tblW','jc','tblCellSpacing','tblInd','tblBorders','shd','tblLayout','tblCellMar','tblLook','tblCaption','tblDescription'],
 'tcPr': ['cnfStyle','tcW','gridSpan','hMerge','vMerge','tcBorders','shd','noWrap','tcMar','textDirection','tcFitText','vAlign','hideMark'],
 'tcBorders': ['top','start','left','bottom','end','right','insideH','insideV','tl2br','tr2bl'],
 'trPr': ['cnfStyle','divId','gridBefore','gridAfter','wBefore','wAfter','cantSplit','trHeight','tblHeader','tblCellSpacing','jc','hidden'],
 'sectPr': ['headerReference','footerReference','footnotePr','endnotePr','type','pgSz','pgMar','paperSrc','pgBorders','lnNumType','pgNumType','cols','formProt','vAlign','noEndnote','titlePg','textDirection','bidi','rtlGutter','docGrid','printerSettings','sectPrChange']}
W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
for tag, order in ORDER.items():
    for el in doc.element.iter(W + tag):
        kids = list(el); kids.sort(key=lambda k: order.index(k.tag[len(W):]) if k.tag.startswith(W) and k.tag[len(W):] in order else 99)
        for k in kids: el.remove(k)
        for k in kids: el.append(k)
doc.core_properties.title = title; doc.core_properties.author = 'Venumadhav Kalluru'
doc.save('CICT348.docx'); print('saved', labels)
